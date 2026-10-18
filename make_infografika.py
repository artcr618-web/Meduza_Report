#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Редактируемая инфографика: карточки можно двигать, менять заливку,
прозрачность, размер и шрифт. Затем — печать."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

NAVY  = RGBColor(4, 30, 60)
DEEP  = RGBColor(3, 22, 48)
WHITE = RGBColor(255, 255, 255)
SKY   = RGBColor(180, 220, 255)
SUN   = RGBColor(255, 210, 63)
CYAN  = RGBColor(53, 208, 255)

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(17.92)
BLANK = prs.slide_layouts[6]

def transparent_card(slide, l, t, w, h, alpha_pct, name):
    """Скруглённая плашка с полупрозрачной заливкой, без обводки."""
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.name = name
    sh.fill.solid()
    sh.fill.fore_color.rgb = NAVY
    spPr = sh._element.find(qn('p:spPr'))
    solid = spPr.find(qn('a:solidFill'))
    clr = solid.find(qn('a:srgbClr'))
    alpha = OxmlElement('a:alpha')
    alpha.set('val', str(int(alpha_pct * 1000)))
    clr.append(alpha)
    sh.line.fill.background()
    sh.adjustments[0] = 0.08
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.08)
    tf.margin_bottom = Inches(0.08)
    return sh

def add_para(tf, text, size, bold=False, color=WHITE, align=PP_ALIGN.LEFT, first=False, space=2):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align; p.space_after = Pt(space); p.space_before = Pt(0)
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = 'Arial'
    return p

# ================= СЛАЙД 1: инфографика =================
s = prs.slides.add_slide(BLANK)
s.shapes.add_picture("doklad_medusy/images/diversity_bg.jpg", Inches(0), Inches(0), Inches(10), Inches(17.92))

title = transparent_card(s, 0.7, 0.4, 8.6, 1.7, 60, "Заголовок — можно двигать")
add_para(title.text_frame, "РАЗНООБРАЗИЕ МЕДУЗ", 44, True, WHITE, PP_ALIGN.CENTER, True, 2)
add_para(title.text_frame, "удивительные стрекающие океана", 24, False, SKY, PP_ALIGN.CENTER)

c1 = transparent_card(s, 5.4, 2.3, 4.1, 3.2, 48, "Карточка 1. Цианея — можно двигать")
tf = c1.text_frame
add_para(tf, "1. Цианея «львиная грива»", 21, True, WHITE, first=True)
add_para(tf, "(Cyanea capillata)", 16, False, SKY)
add_para(tf, "Самая большая медуза: купол до 2,5 м, щупальца до 37 м!", 18)

c2 = transparent_card(s, 5.3, 12.4, 4.2, 2.8, 48, "Карточка 2. Кубомедуза — можно двигать")
tf = c2.text_frame
add_para(tf, "2. Кубомедуза", 21, True, WHITE, first=True)
add_para(tf, "(Chironex fleckeri)", 16, False, SKY)
add_para(tf, "«Морская оса» — самая ядовитая, 24 глаза!", 18)

c3 = transparent_card(s, 2.4, 14.8, 4.7, 2.5, 48, "Карточка 3. Бессмертная — можно двигать")
tf = c3.text_frame
add_para(tf, "3. Бессмертная медуза", 21, True, WHITE, first=True)
add_para(tf, "(Turritopsis dohrnii)", 16, False, SKY)
add_para(tf, "Крошка 5 мм умеет молодеть и жить заново!", 18)

# ================= СЛАЙД 2: инструкция =================
s2 = prs.slides.add_slide(BLANK)
fill = s2.background.fill; fill.solid(); fill.fore_color.rgb = DEEP
t = s2.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(8.6), Inches(1.0))
add_para(t.text_frame, "Как двигать карточки 🎨", 36, True, SUN, PP_ALIGN.LEFT, True)
tips = [
    ("Двигать: ", "кликни по карточке левой кнопкой мыши и тяни в любое место картинки."),
    ("Размер: ", "потяни за белые кружочки по углам карточки."),
    ("Заливка и прозрачность: ", "правый клик по карточке → «Формат фигуры» → «Заливка» → ползунок «Прозрачность»."),
    ("Шрифт: ", "выдели текст → вкладка «Главная» → выбери шрифт, размер и цвет."),
    ("Печать: ", "«Файл» → «Печать». Картинка вертикальная — хорошо ляжет на лист А4."),
    ("Найти карточку: ", "вкладка «Главная» → «Выделить» → «Область выделения» — там все карточки подписаны."),
]
y = 1.7
for b, rest in tips:
    sh = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(y), Inches(8.6), Inches(1.15))
    sh.fill.solid(); sh.fill.fore_color.rgb = NAVY
    sh.line.color.rgb = CYAN; sh.line.width = Pt(2)
    sh.adjustments[0] = 0.1
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.25)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = b; r.font.size = Pt(19); r.font.bold = True; r.font.color.rgb = SUN; r.font.name = 'Arial'
    r2 = p.add_run(); r2.text = rest; r2.font.size = Pt(19); r2.font.color.rgb = WHITE; r2.font.name = 'Arial'
    y += 1.32

prs.save("doklad_medusy/Infografika_Raznoobrazie_Meduz.pptx")
print("infografika saved")
